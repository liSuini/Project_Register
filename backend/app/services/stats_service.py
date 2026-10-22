from io import BytesIO

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.entities import AcceptanceRecord, Project, PurchaseOrder
from app.schemas.dtos import CaliberItem, StatsOut


def calibers(db: Session) -> StatsOut:
    """三口径聚合（ADR-0004：采购额/回款额只由明细聚合，不冗余到项目）。"""
    fwd_count = db.query(func.count(Project.id)) \
        .filter(Project.contract_amount > 0).scalar()
    fwd_amount = db.query(func.coalesce(func.sum(Project.contract_amount), 0)) \
        .filter(Project.contract_amount > 0).scalar()

    back = db.query(func.count(func.distinct(PurchaseOrder.project_id)),
                    func.coalesce(func.sum(PurchaseOrder.amount), 0)).one()
    deliv = db.query(func.count(func.distinct(AcceptanceRecord.project_id)),
                     func.coalesce(func.sum(AcceptanceRecord.received_amount), 0)).one()
    return StatsOut(
        forward=CaliberItem(count=int(fwd_count), amount=float(fwd_amount)),
        backward=CaliberItem(count=int(back[0]), amount=float(back[1])),
        deliver=CaliberItem(count=int(deliv[0]), amount=float(deliv[1])))


def export_xlsx(db: Session) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "前向-签约明细"
    ws1.append(["项目名称", "客户", "合同编号", "当前阶段", "签约金额(元)",
                "负责人", "起始日期", "结束日期"])
    for p in db.query(Project).order_by(Project.id).all():
        ws1.append([p.name, p.customer, p.contract_no, p.current_stage,
                    float(p.contract_amount or 0), p.owner,
                    str(p.start_date or ""), str(p.end_date or "")])

    ws2 = wb.create_sheet("后向-采购明细")
    ws2.append(["项目名称", "采购名称", "供应商", "采购金额(元)", "采购日期", "备注"])
    for r in db.query(PurchaseOrder).order_by(PurchaseOrder.id).all():
        project = db.get(Project, r.project_id)
        ws2.append([project.name if project else "", r.name, r.supplier,
                    float(r.amount or 0), str(r.order_date or ""), r.note])

    ws3 = wb.create_sheet("交付-回款明细")
    ws3.append(["项目名称", "验收名称", "回款金额(元)", "验收日期", "备注"])
    for r in db.query(AcceptanceRecord).order_by(AcceptanceRecord.id).all():
        project = db.get(Project, r.project_id)
        ws3.append([project.name if project else "", r.name,
                    float(r.received_amount or 0),
                    str(r.acceptance_date or ""), r.note])

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
