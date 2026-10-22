def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200
    assert r.json() == {"status": "ok"}


def test_full_chain(client):
    """一账全链路：建项目→锚点→子环节→附件→采购验收→检索→统计→导出。"""
    from io import BytesIO

    from openpyxl import load_workbook

    # 1) 建项目（存量从采购起步）
    r = client.post("/api/projects", json={
        "name": "政务云灾备扩容", "customer": "省政务服务中心",
        "initial_stage": "采购", "contract_amount": 8600000})
    assert r.status_code == 201, r.text
    pid = r.json()["id"]

    # 2) 一账详情：9 锚点
    ledger = client.get(f"/api/projects/{pid}/ledger").json()
    assert len(ledger["tree"]) == 9
    assert ledger["project"]["current_stage"] == "采购"
    assert [t["name"] for t in ledger["tree"]][4] == "采购"

    # 3) 阶段流转 + 非法阶段
    assert client.patch(f"/api/projects/{pid}/stage",
                        json={"stage": "实施"}).status_code == 200
    assert client.patch(f"/api/projects/{pid}/stage",
                        json={"stage": "不存在"}).status_code == 422

    # 4) 子环节 + 嵌套
    anchor_id = ledger["tree"][4]["id"]            # 采购锚点
    r = client.post("/api/nodes", json={"parent_id": anchor_id, "name": "硬件采购"})
    assert r.status_code == 201
    sub_id = r.json()["id"]
    assert client.post("/api/nodes",
                       json={"parent_id": sub_id, "name": "二标段"}).status_code == 201

    # 5) 附件上传/下载/列表/删除
    r = client.post(f"/api/nodes/{sub_id}/attachments",
                    files={"file": ("验收报告.pdf", BytesIO(b"%PDF-fake"),
                                    "application/pdf")})
    assert r.status_code == 201, r.text
    att = r.json()
    assert att["original_name"] == "验收报告.pdf"
    dl = client.get(f"/api/attachments/{att['id']}/download")
    assert dl.status_code == 200 and dl.content == b"%PDF-fake"
    assert len(client.get(f"/api/nodes/{sub_id}/attachments").json()) == 1
    assert client.delete(f"/api/attachments/{att['id']}").status_code == 204

    # 6) 非空节点删除保护（子环节下还有"二标段"）
    r = client.delete(f"/api/nodes/{sub_id}")
    assert r.status_code == 409
    assert r.json()["code"] == "conflict"

    # 7) 采购/验收明细
    assert client.post(f"/api/projects/{pid}/purchases",
                       json={"name": "防火墙", "supplier": "华为",
                             "amount": 1200000}).status_code == 201
    assert client.post(f"/api/projects/{pid}/acceptances",
                       json={"name": "初验",
                            "received_amount": 300000}).status_code == 201
    ledger = client.get(f"/api/projects/{pid}/ledger").json()
    assert float(ledger["totals"]["purchase_amount"]) == 1200000
    assert float(ledger["totals"]["received_amount"]) == 300000

    # 8) 检索
    assert len(client.get("/api/projects",
                          params={"keyword": "灾备"}).json()) == 1
    assert len(client.get("/api/projects",
                          params={"stage": "实施"}).json()) == 1

    # 9) 统计 + 导出可打开
    stats = client.get("/api/stats").json()
    assert stats["forward"]["amount"] == 8600000
    assert stats["backward"]["amount"] == 1200000
    exp = client.get("/api/stats/export")
    assert exp.status_code == 200
    wb = load_workbook(BytesIO(exp.content))
    assert wb.sheetnames == ["前向-签约明细", "后向-采购明细", "交付-回款明细"]


def test_404_contract(client):
    assert client.get("/api/projects/999/ledger").status_code == 404
    assert client.get("/api/projects/999/ledger").json()["code"] == "not_found"


def test_app_error_contract(client):
    """领域异常统一错误契约：{code, message}，409 映射。"""
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.core.config import make_settings
    from app.errors import ConflictError
    from app.main import create_app

    app = create_app(make_settings(str(client.app.state.settings.data_dir) + "-err"))

    def boom():
        raise ConflictError("测试冲突")

    app.add_api_route("/api/boom", boom, methods=["GET"])
    with TestClient(app, raise_server_exceptions=False) as c:
        r = c.get("/api/boom")
        assert r.status_code == 409
        assert r.json() == {"code": "conflict", "message": "测试冲突"}
