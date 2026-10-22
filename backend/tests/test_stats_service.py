from io import BytesIO

from app.schemas.dtos import AcceptanceCreate, ProjectCreate, PurchaseCreate
from app.services import detail_service as ds
from app.services import project_service as ps


def seed(db_session):
    p1 = ps.create_project(db_session, ProjectCreate(
        name="甲", initial_stage="采购", contract_amount=8600000))
    ps.create_project(db_session, ProjectCreate(name="乙", initial_stage="商机",
                                                contract_amount=0))
    ps.create_project(db_session, ProjectCreate(name="丙", initial_stage="实施",
                                                contract_amount=5000000))
    ds.create_purchase(db_session, p1.id, PurchaseCreate(name="a", amount=1200000))
    ds.create_purchase(db_session, p1.id, PurchaseCreate(name="b", amount=800000))
    ds.create_acceptance(db_session, p1.id,
                         AcceptanceCreate(name="初验", received_amount=300000))
    return p1


def test_calibers(db_session):
    from app.services.stats_service import calibers

    seed(db_session)
    s = calibers(db_session)
    assert (s.forward.count, s.forward.amount) == (2, 13600000)     # 甲+丙有签约额
    assert (s.backward.count, s.backward.amount) == (1, 2000000)
    assert (s.deliver.count, s.deliver.amount) == (1, 300000)


def test_calibers_empty(db_session):
    from app.services.stats_service import calibers

    s = calibers(db_session)
    assert (s.forward.count, s.forward.amount) == (0, 0)
    assert (s.backward.count, s.backward.amount) == (0, 0)
    assert (s.deliver.count, s.deliver.amount) == (0, 0)


def test_export_xlsx_openable(db_session):
    from openpyxl import load_workbook

    from app.services.stats_service import export_xlsx

    seed(db_session)
    content = export_xlsx(db_session)
    wb = load_workbook(BytesIO(content))
    assert wb.sheetnames == ["前向-签约明细", "后向-采购明细", "交付-回款明细"]
    ws = wb["后向-采购明细"]
    assert ws.max_row >= 3          # 表头 + 2 笔明细
    assert ws.cell(1, 1).value == "项目名称"
    fwd = wb["前向-签约明细"]
    assert fwd.cell(2, 1).value == "甲"   # 明细含项目名
